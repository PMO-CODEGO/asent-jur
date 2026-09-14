import re
import unicodedata
from datetime import datetime, date

import openpyxl


STOPWORDS_SIGLA = {'DE', 'DA', 'DO', 'DOS', 'DAS', 'E', 'A', 'O', 'EM'}

# (índice da coluna na planilha) -> (campo, palavras-chave para validar o cabeçalho)
# Modelo fixo da planilha "Cadastro de Módulos" (mesma estrutura para qualquer distrito/município;
# reproduz o tratamento do notebook CADASTRO_DE_MÓDULOS.ipynb usado pela equipe).
COLUNAS_MODELO = [
    ('_municipio_raw', ['MUNICIPIO']),
    ('matricula_loteamento', ['MATRICULA', 'LOTEAMENTO']),
    ('distrito', ['DISTRITO']),
    ('_sigla_raw', ['SIGLA']),
    ('quadra', ['QUADRA']),
    ('qtd_modulos', ['MODULO']),
    ('logradouro', ['LOGRADOURO']),
    ('area_lote_m2', ['AREA', 'LOTE']),
    ('matricula_modulo', ['MATRICULA', 'MODULO']),
    ('cci', ['CCI']),
    ('inscricao_municipal', ['INSCRICAO']),
    ('area_institucional', ['AREA', 'INSTITUCIONAL']),
    ('empresa', ['EMPRESA']),
    ('cnpj', ['CNPJ']),
    ('nome_representante_legal', ['REPRESENTANTE']),
    ('telefone_representante_legal', ['TELEFONE']),
    ('email_representante_legal', ['MAIL']),
    ('processo_sei', ['PROCESSO']),
    ('ramo_de_atividade', ['RAMO']),
    ('status_de_assentamento', ['SITUACAO', 'ASSENTAMENTO']),
    ('data_escrituracao', ['DATA', 'ESCRITURA']),
    ('registro_na_matricula', ['REGISTRO', 'MATRICULA']),
    ('empresa_anterior', ['NOME', 'EMPRESA']),
    ('cnpj_anterior', ['CNPJ']),
    ('processo_anterior', ['NUMERO', 'PROCESSO']),
    ('relatorio_vistoria', ['RELATORIO', 'VISTORIA']),
    ('ultima_vistoria', ['VISTORIA']),
    ('taxa_ocupacao_imovel', ['TAXA', 'OCUPACAO']),
    ('atividade_industrial', ['ATIVIDADE', 'INDUSTRIAL']),
    ('irregularidades', ['IRREGULARIDADE']),
    ('empregos_gerados', ['EMPREGO']),
    ('projeto_ocupacao_area', ['PROJETO']),
    ('data_aprovacao_poa', ['APROVACAO']),
    ('cronograma_fisico_obra_meses', ['CRONOGRAMA']),
    ('imovel_regular_irregular', ['REGULAR', 'IRREGULAR']),
    ('observacoes', ['OBSERVA']),
]

CAMPOS_INTEIROS = {'matricula_loteamento', 'quadra', 'qtd_modulos', 'matricula_modulo', 'empregos_gerados', 'cronograma_fisico_obra_meses'}
CAMPOS_FLOAT = {'area_lote_m2', 'taxa_ocupacao_imovel'}
CAMPOS_DATA = {'data_escrituracao', 'ultima_vistoria', 'data_aprovacao_poa'}


def _remover_acentos(texto):
    nfkd = unicodedata.normalize('NFKD', str(texto))
    return ''.join(c for c in nfkd if not unicodedata.combining(c))


def _normalizar_cabecalho(valor):
    if valor is None:
        return ''
    texto = str(valor).replace('\n', ' ')
    texto = re.sub(r'\s+', ' ', texto).strip().upper()
    return _remover_acentos(texto)


def _normalizar_valor_texto(valor):
    if valor is None:
        return None
    texto = re.sub(r'\s+', ' ', str(valor)).strip()
    if texto in ('', '-'):
        return None
    return _remover_acentos(texto).upper()


def _parse_inteiro(valor):
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return int(valor)
    texto = str(valor).strip()
    if texto in ('', '-'):
        return None
    try:
        return int(float(texto.replace(',', '.')))
    except ValueError:
        return None


def _parse_float(valor):
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return round(float(valor), 4)
    texto = str(valor).strip()
    if texto in ('', '-'):
        return None
    if re.match(r'^-?\d{1,3}(\.\d{3})*(,\d+)?$', texto):
        texto = texto.replace('.', '').replace(',', '.')
    else:
        texto = texto.replace(',', '.')
    try:
        return round(float(texto), 4)
    except ValueError:
        return None


def _parse_data(valor):
    if valor is None:
        return None
    if isinstance(valor, (datetime, date)):
        return valor.strftime('%Y-%m-%d')
    texto = str(valor).strip()
    if texto in ('', '-'):
        return None
    for formato in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y'):
        try:
            return datetime.strptime(texto, formato).strftime('%Y-%m-%d')
        except ValueError:
            continue
    return None


def _gerar_sigla(sigla_raw, distrito):
    sigla_normalizada = _normalizar_valor_texto(sigla_raw)
    if sigla_normalizada:
        return sigla_normalizada
    if not distrito:
        return None
    palavras = distrito.split()
    iniciais = [p[0] for p in palavras if p not in STOPWORDS_SIGLA and p]
    return ''.join(iniciais).upper() or None


def gerar_id_modulo(sigla, quadra, modulo, logradouro):
    sigla_limpa = re.sub(r'[\-\s]+', '', str(sigla or ''))
    quadra_limpa = re.sub(r'[\-\s]+', '', str(quadra)) if quadra not in (None, '', '-') else ''
    modulo_limpo = re.sub(r'[\-\s]+', '', str(modulo)) if modulo not in (None, '', '-') else ''

    if not quadra_limpa or not modulo_limpo:
        logradouro_limpo = re.sub(r'[\-\s]+', '', str(logradouro or ''))
        return _remover_acentos(f'{sigla_limpa}{logradouro_limpo}').upper()

    return f'{sigla_limpa}Q{quadra_limpa}M{modulo_limpo}'


def listar_abas_candidatas(caminho_arquivo):
    wb = openpyxl.load_workbook(caminho_arquivo, read_only=True)
    nomes = wb.sheetnames
    wb.close()
    return nomes


def detectar_aba_padrao(nomes_abas):
    candidatas = [n for n in nomes_abas if 'MAPA' not in _normalizar_cabecalho(n)]
    if not candidatas:
        return nomes_abas[0] if nomes_abas else None
    return candidatas[0]


def _validar_cabecalho(header_row):
    if len(header_row) < len(COLUNAS_MODELO):
        raise ValueError(
            f'A aba selecionada tem apenas {len(header_row)} colunas; eram esperadas pelo menos '
            f'{len(COLUNAS_MODELO)}, no padrão da planilha de Cadastro de Módulos. Confirme se é a aba correta.'
        )
    for idx, (campo, palavras_chave) in enumerate(COLUNAS_MODELO):
        cabecalho_norm = _normalizar_cabecalho(header_row[idx])
        if not all(palavra in cabecalho_norm for palavra in palavras_chave):
            raise ValueError(
                f'A coluna {idx + 1} da aba selecionada tem o cabeçalho "{header_row[idx]}", que não '
                f'corresponde ao esperado para esse modelo de planilha (Cadastro de Módulos). Confirme se é a aba correta.'
            )


def processar_planilha_cadastro(caminho_arquivo, nome_aba):
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    if nome_aba not in wb.sheetnames:
        raise ValueError(f"A aba '{nome_aba}' não foi encontrada na planilha.")
    ws = wb[nome_aba]

    linhas = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    wb.close()

    if len(linhas) < 4:
        raise ValueError('A planilha não tem linhas suficientes para conter os dados esperados (cabeçalho na linha 2).')

    header_row = linhas[1]
    _validar_cabecalho(header_row)

    registros = []
    for row in linhas[3:]:
        if row is None:
            continue

        valores = {}
        for idx, (campo, _palavras_chave) in enumerate(COLUNAS_MODELO):
            valores[campo] = row[idx] if idx < len(row) else None

        matricula_modulo = _parse_inteiro(valores['matricula_modulo'])
        if matricula_modulo is None:
            # linhas sem matrícula de módulo são totais/infraestrutura (ex: "TOTAL DA GLEBA",
            # "Vias Públicas"), não módulos cadastrados individualmente.
            continue

        distrito = _normalizar_valor_texto(valores['distrito'])
        sigla = _gerar_sigla(valores['_sigla_raw'], distrito)
        quadra = _parse_inteiro(valores['quadra'])
        qtd_modulos = _parse_inteiro(valores['qtd_modulos'])
        logradouro = _normalizar_valor_texto(valores['logradouro'])

        registro = {
            'municipio': _normalizar_valor_texto(valores['_municipio_raw']),
            'matricula_loteamento': _parse_inteiro(valores['matricula_loteamento']),
            'distrito': distrito,
            'sigla_loteamento': sigla,
            'quadra': quadra,
            'qtd_modulos': qtd_modulos,
            'logradouro': logradouro,
            'area_lote_m2': _parse_float(valores['area_lote_m2']),
            'matricula_modulo': matricula_modulo,
            'cci': _normalizar_valor_texto(valores['cci']),
            'inscricao_municipal': _normalizar_valor_texto(valores['inscricao_municipal']),
            'area_institucional': _normalizar_valor_texto(valores['area_institucional']),
            'empresa': _normalizar_valor_texto(valores['empresa']),
            'cnpj': _normalizar_valor_texto(valores['cnpj']),
            'nome_representante_legal': _normalizar_valor_texto(valores['nome_representante_legal']),
            'telefone_representante_legal': _normalizar_valor_texto(valores['telefone_representante_legal']),
            'email_representante_legal': _normalizar_valor_texto(valores['email_representante_legal']),
            'processo_sei': _normalizar_valor_texto(valores['processo_sei']),
            'ramo_de_atividade': _normalizar_valor_texto(valores['ramo_de_atividade']),
            'status_de_assentamento': _normalizar_valor_texto(valores['status_de_assentamento']),
            'data_escrituracao': _parse_data(valores['data_escrituracao']),
            'data_contrato_de_compra_e_venda': None,
            'registro_na_matricula': _normalizar_valor_texto(valores['registro_na_matricula']),
            'empresa_anterior': _normalizar_valor_texto(valores['empresa_anterior']),
            'cnpj_anterior': _normalizar_valor_texto(valores['cnpj_anterior']),
            'processo_anterior': _normalizar_valor_texto(valores['processo_anterior']),
            'relatorio_vistoria': _normalizar_valor_texto(valores['relatorio_vistoria']),
            'ultima_vistoria': _parse_data(valores['ultima_vistoria']),
            'taxa_ocupacao_imovel': _parse_float(valores['taxa_ocupacao_imovel']),
            'atividade_industrial': _normalizar_valor_texto(valores['atividade_industrial']),
            'irregularidades': _normalizar_valor_texto(valores['irregularidades']),
            'empregos_gerados': _parse_inteiro(valores['empregos_gerados']),
            'projeto_ocupacao_area': _normalizar_valor_texto(valores['projeto_ocupacao_area']),
            'data_aprovacao_poa': _parse_data(valores['data_aprovacao_poa']),
            'cronograma_fisico_obra_meses': _parse_inteiro(valores['cronograma_fisico_obra_meses']),
            'imovel_regular_irregular': _normalizar_valor_texto(valores['imovel_regular_irregular']),
            'observacoes': _normalizar_valor_texto(valores['observacoes']),
        }
        registro['id_modulo'] = gerar_id_modulo(sigla, quadra, qtd_modulos, logradouro)

        registros.append(registro)

    return registros
